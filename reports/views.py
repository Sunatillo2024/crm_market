import csv
from datetime import timedelta
from decimal import Decimal

from django.db.models import (
    Avg,
    Count,
    DecimalField,
    ExpressionWrapper,
    F,
    Sum,
)
from django.db.models.functions import TruncDate
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.views.decorators.http import require_GET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from accounts.decorators import roles_required
from accounts.models import User
from sales.models import Sale, SaleItem


def get_report_period(request):
    """
    GET-so'rovdan 'period', 'date_from', 'date_to' parametrlarini olib,
    hisobot davrini (date_from, date_to, selected_period) qaytaradi.
    """
    today = timezone.localdate()

    selected_period = request.GET.get("period", "month").strip()

    valid_periods = {"today", "7days", "month", "custom"}
    if selected_period not in valid_periods:
        selected_period = "month"

    if selected_period == "today":
        date_from = today
        date_to = today

    elif selected_period == "7days":
        date_from = today - timedelta(days=6)
        date_to = today

    elif selected_period == "month":
        date_from = today.replace(day=1)
        date_to = today

    else:  # custom
        raw_date_from = request.GET.get("date_from", "").strip()
        raw_date_to = request.GET.get("date_to", "").strip()

        date_from = parse_date(raw_date_from)
        date_to = parse_date(raw_date_to)

        if date_from is None:
            date_from = today.replace(day=1)

        if date_to is None:
            date_to = today

    if date_from > date_to:
        date_from, date_to = date_to, date_from

    return date_from, date_to, selected_period


def get_profit_expression():
    return ExpressionWrapper(
        (F("unit_price") - F("purchase_price")) * F("quantity"),
        output_field=DecimalField(
            max_digits=24,
            decimal_places=4,
        ),
    )


@roles_required(User.Role.OWNER)
def report_index(request):
    store = request.user.store

    # 3 ta qiymatni qabul qilib olamiz
    date_from, date_to, selected_period = get_report_period(request)

    period_sales = Sale.objects.filter(
        store=store,
        created_at__date__range=(
            date_from,
            date_to,
        ),
    )

    completed_sales = period_sales.filter(
        status=Sale.Status.COMPLETED,
    )

    completed_summary = completed_sales.aggregate(
        revenue=Sum("total"),
        average_check=Avg("total"),
    )

    revenue = completed_summary["revenue"] or Decimal("0")
    average_check = completed_summary["average_check"] or Decimal("0")

    sale_items = SaleItem.objects.filter(
        sale__in=completed_sales,
    )

    total_profit = sale_items.aggregate(
        total=Sum(get_profit_expression()),
    )["total"] or Decimal("0")

    if revenue > 0:
        profit_margin = total_profit / revenue * Decimal("100")
    else:
        profit_margin = Decimal("0")

    cash_sales = completed_sales.filter(
        payment_method=Sale.PaymentMethod.CASH,
    )

    card_sales = completed_sales.filter(
        payment_method=Sale.PaymentMethod.CARD,
    )

    cash_total = cash_sales.aggregate(total=Sum("total"))["total"] or Decimal("0")
    card_total = card_sales.aggregate(total=Sum("total"))["total"] or Decimal("0")

    # Kunlik savdolar
    daily_rows = list(
        completed_sales.order_by()
        .annotate(
            day=TruncDate(
                "created_at",
                tzinfo=timezone.get_current_timezone(),
            ),
        )
        .values("day")
        .annotate(
            revenue=Sum("total"),
            sales_count=Count("pk"),
        )
        .order_by("day")
    )

    daily_values = {row["day"]: row for row in daily_rows}

    daily_report = []
    current_day = date_from

    while current_day <= date_to:
        row = daily_values.get(current_day, {})

        daily_report.append(
            {
                "date": current_day,
                "revenue": row.get("revenue") or Decimal("0"),
                "sales_count": row.get("sales_count") or 0,
            }
        )

        current_day += timedelta(days=1)

    max_daily_revenue = max(
        (row["revenue"] for row in daily_report),
        default=Decimal("0"),
    )

    for row in daily_report:
        if max_daily_revenue > 0 and row["revenue"] > 0:
            height = float(row["revenue"] / max_daily_revenue * Decimal("100"))
            row["bar_height"] = f"{max(height, 4):.2f}"
        else:
            row["bar_height"] = "0"

    # Eng ko‘p sotilgan mahsulotlar
    top_products = list(
        sale_items.order_by()
        .values(
            "product_id",
            "product_name",
            "unit",
        )
        .annotate(
            quantity_sold=Sum("quantity"),
            revenue=Sum("line_total"),
            profit=Sum(get_profit_expression()),
        )
        .order_by(
            "-revenue",
            "product_name",
        )[:10]
    )

    unit_labels = dict(SaleItem._meta.get_field("unit").flatchoices)

    for product in top_products:
        product["unit_label"] = unit_labels.get(
            product["unit"],
            product["unit"],
        )

    # Kassirlar statistikasi
    cashier_stats = list(
        completed_sales.order_by()
        .values("cashier_id")
        .annotate(
            sales_count=Count("pk"),
            revenue=Sum("total"),
            average_check=Avg("total"),
        )
        .order_by(
            "-revenue",
            "-sales_count",
        )[:10]
    )

    cashier_ids = [
        row["cashier_id"] for row in cashier_stats if row["cashier_id"] is not None
    ]

    cashier_map = User.objects.in_bulk(cashier_ids)

    for row in cashier_stats:
        cashier = cashier_map.get(row["cashier_id"])
        row["cashier_name"] = str(cashier) if cashier else "O‘chirilgan kassir"

    context = {
        "selected_period": selected_period,  # Shablonga (HTML) tanlangan davr tugmasini faol ko'rsatish uchun uzatildi
        "date_from": date_from,
        "date_to": date_to,
        "sales_count": completed_sales.count(),
        "revenue": revenue,
        "average_check": average_check,
        "total_profit": total_profit,
        "profit_margin": profit_margin,
        "cash_count": cash_sales.count(),
        "cash_total": cash_total,
        "card_count": card_sales.count(),
        "card_total": card_total,
        "cancelled_count": period_sales.filter(
            status=Sale.Status.CANCELLED,
        ).count(),
        "daily_report": daily_report,
        "top_products": top_products,
        "cashier_stats": cashier_stats,
    }

    return render(
        request,
        "reports/index.html",
        context,
    )


def get_export_items(
    *,
    store,
    date_from,
    date_to,
):
    return (
        SaleItem.objects.select_related(
            "sale",
            "sale__cashier",
        )
        .filter(
            sale__store=store,
            sale__status=Sale.Status.COMPLETED,
            sale__created_at__date__range=(
                date_from,
                date_to,
            ),
        )
        .order_by(
            "sale__created_at",
            "sale_id",
            "pk",
        )
    )


def spreadsheet_safe_text(value):
    text = str(
        value or "",
    )

    if text.startswith(
        (
            "=",
            "+",
            "-",
            "@",
        )
    ):
        return f"'{text}"

    return text


@roles_required(User.Role.OWNER)
@require_GET
def report_export_csv(request):
    (
        date_from,
        date_to,
        selected_period,
    ) = get_report_period(
        request,
    )

    items = get_export_items(
        store=request.user.store,
        date_from=date_from,
        date_to=date_to,
    )

    filename = (
        f"sales-report-"
        f"{date_from.isoformat()}-"
        f"{date_to.isoformat()}.csv"
    )

    response = HttpResponse(
        content_type=(
            "text/csv; charset=utf-8"
        ),
    )

    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{filename}"'
    )

    # Excel o‘zbekcha matnlarni to‘g‘ri ochishi uchun
    response.write(
        "\ufeff",
    )

    writer = csv.writer(
        response,
    )

    writer.writerow(
        [
            "Savdo raqami",
            "Sana",
            "Kassir",
            "Mahsulot",
            "Barkod",
            "Miqdor",
            "Birlik",
            "Sotuv narxi",
            "Olish narxi",
            "Jami",
            "Foyda",
            "To‘lov turi",
        ]
    )

    for item in items:
        sale = item.sale

        created_at = timezone.localtime(
            sale.created_at,
        )

        writer.writerow(
            [
                spreadsheet_safe_text(
                    sale.sale_number,
                ),
                created_at.strftime(
                    "%d.%m.%Y %H:%M",
                ),
                spreadsheet_safe_text(
                    sale.cashier
                    if sale.cashier
                    else "O‘chirilgan kassir",
                ),
                spreadsheet_safe_text(
                    item.product_name,
                ),
                spreadsheet_safe_text(
                    item.barcode,
                ),
                item.quantity,
                spreadsheet_safe_text(
                    item.get_unit_display(),
                ),
                item.unit_price,
                item.purchase_price,
                item.line_total,
                item.profit,
                spreadsheet_safe_text(
                    sale.get_payment_method_display(),
                ),
            ]
        )

    return response


@roles_required(User.Role.OWNER)
@require_GET
def report_export_excel(request):
    (
        date_from,
        date_to,
        selected_period,
    ) = get_report_period(
        request,
    )

    items = get_export_items(
        store=request.user.store,
        date_from=date_from,
        date_to=date_to,
    )

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Savdolar"

    headers = [
        "Savdo raqami",
        "Sana",
        "Kassir",
        "Mahsulot",
        "Barkod",
        "Miqdor",
        "Birlik",
        "Sotuv narxi",
        "Olish narxi",
        "Jami",
        "Foyda",
        "To‘lov turi",
    ]

    worksheet.append(
        headers,
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="3763F4",
    )

    for cell in worksheet[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for item in items:
        sale = item.sale

        created_at = timezone.localtime(
            sale.created_at,
        ).replace(
            tzinfo=None,
        )

        worksheet.append(
            [
                spreadsheet_safe_text(
                    sale.sale_number,
                ),
                created_at,
                spreadsheet_safe_text(
                    sale.cashier
                    if sale.cashier
                    else "O‘chirilgan kassir",
                ),
                spreadsheet_safe_text(
                    item.product_name,
                ),
                spreadsheet_safe_text(
                    item.barcode,
                ),
                float(
                    item.quantity,
                ),
                spreadsheet_safe_text(
                    item.get_unit_display(),
                ),
                float(
                    item.unit_price,
                ),
                float(
                    item.purchase_price,
                ),
                float(
                    item.line_total,
                ),
                float(
                    item.profit,
                ),
                spreadsheet_safe_text(
                    sale.get_payment_method_display(),
                ),
            ]
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        worksheet.cell(
            row=row_number,
            column=2,
        ).number_format = (
            "dd.mm.yyyy hh:mm"
        )

        for column_number in {
            6,
            8,
            9,
            10,
            11,
        }:
            worksheet.cell(
                row=row_number,
                column=column_number,
            ).number_format = (
                '#,##0.00'
            )

    column_widths = [
        20,
        19,
        22,
        32,
        18,
        13,
        14,
        16,
        16,
        16,
        16,
        18,
    ]

    for index, width in enumerate(
        column_widths,
        start=1,
    ):
        worksheet.column_dimensions[
            get_column_letter(index)
        ].width = width

    filename = (
        f"sales-report-"
        f"{date_from.isoformat()}-"
        f"{date_to.isoformat()}.xlsx"
    )

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
    )

    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{filename}"'
    )

    workbook.save(
        response,
    )

    return response